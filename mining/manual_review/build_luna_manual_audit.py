#!/usr/bin/env python3
"""Build a blinded Excel workbook for the GPT-Luna manual audit."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path

import pandas as pd
import pyarrow.parquet as pq
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


SEED = "luna-manual-audit-v1-20260924"
ILLEGAL_XML = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--task-decisions", type=Path, required=True)
    parser.add_argument("--curated-labels", type=Path, required=True)
    parser.add_argument("--balanced-sample", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--source-revision", default="unknown")
    return parser.parse_args()


def clean_cell(value: object, limit: int = 30_000) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    if isinstance(value, (list, tuple)):
        text = "\n".join(str(item) for item in value)
    elif isinstance(value, dict):
        text = json.dumps(value, ensure_ascii=False, sort_keys=True)
    else:
        text = str(value)
    text = ILLEGAL_XML.sub("", text)
    if len(text) > limit:
        text = text[: limit - 22] + "\n[TRUNCATED IN EXCEL]"
    if text.startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text


def stable_hash(seed: str, *values: object) -> str:
    payload = "|".join([seed, *(str(value) for value in values)])
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def key_frame(frame: pd.DataFrame) -> pd.Series:
    return frame["repo_id"].astype(str) + ":" + frame["number"].astype(str)


def choose(frame: pd.DataFrame, stratum: str, quota: int) -> pd.DataFrame:
    selected = frame.copy()
    selected["audit_stratum"] = stratum
    selected["selection_hash"] = [
        stable_hash(SEED, stratum, repo_id, number)
        for repo_id, number in zip(selected["repo_id"], selected["number"], strict=True)
    ]
    selected = selected.sort_values("selection_hash", kind="stable")
    if len(selected) < quota:
        raise ValueError(f"Stratum {stratum!r} has {len(selected)} rows; needs {quota}")
    population = len(selected)
    selected = selected.head(quota).copy()
    selected["stratum_population"] = population
    selected["stratum_sample_size"] = quota
    selected["inclusion_probability"] = quota / population
    selected["sampling_weight"] = population / quota
    return selected


def build_sample(
    curated_path: Path, task_path: Path, balanced_path: Path
) -> tuple[pd.DataFrame, pd.DataFrame]:
    curated_columns = [
        "repo_id",
        "repo_full_name",
        "number",
        "html_url",
        "aidev_attribution_label",
        "aidev_task_type",
        "aidev_task_type_reason",
        "aidev_task_type_confidence",
        "aidev_task_type_method",
        "aidev_task_type_status",
        "aidev_task_type_model",
        "aidev_task_type_classifier",
    ]
    curated = pq.read_table(curated_path, columns=curated_columns).to_pandas()

    availability = pq.read_table(
        task_path, columns=["repo_id", "number", "deleted_repo"]
    ).to_pandas()
    curated = curated.merge(
        availability, on=["repo_id", "number"], how="left", validate="one_to_one"
    )
    curated = curated[
        curated["aidev_task_type_method"].eq("llm")
        & curated["aidev_task_type_status"].eq("classified")
        & curated["aidev_attribution_label"].isin(["agentic", "human_candidate"])
        & ~curated["deleted_repo"].fillna(False)
        & curated["html_url"].fillna("").ne("")
    ].copy()
    curated["population_key"] = key_frame(curated)

    balanced = pq.read_table(
        balanced_path,
        columns=["repo_id", "number", "aidev_task_type_method"],
    ).to_pandas()
    final_keys = set(
        key_frame(balanced[balanced["aidev_task_type_method"].eq("llm")])
    )
    curated["in_final_sample"] = curated["population_key"].isin(final_keys)
    curated["audit_arm"] = curated["aidev_attribution_label"]
    curated["negative_label_group"] = curated["aidev_task_type"].where(
        curated["aidev_task_type"].isin(["fix", "refactor", "feat"]), "other"
    )

    selected_parts: list[pd.DataFrame] = []
    positive = curated[curated["aidev_task_type"].eq("perf")]
    negative = curated[~curated["aidev_task_type"].eq("perf")]

    for arm in ["agentic", "human_candidate"]:
        selected_parts.append(
            choose(
                positive[positive["in_final_sample"] & positive["audit_arm"].eq(arm)],
                f"positive_final_{arm}",
                15,
            )
        )
        selected_parts.append(
            choose(
                positive[~positive["in_final_sample"] & positive["audit_arm"].eq(arm)],
                f"positive_outside_{arm}",
                15,
            )
        )
        for label_group in ["fix", "refactor", "feat", "other"]:
            selected_parts.append(
                choose(
                    negative[
                        negative["audit_arm"].eq(arm)
                        & negative["negative_label_group"].eq(label_group)
                    ],
                    f"negative_{arm}_{label_group}",
                    5,
                )
            )

    selected = pd.concat(selected_parts, ignore_index=True)
    if len(selected) != 100 or selected["population_key"].duplicated().any():
        raise AssertionError("The audit sample must contain 100 unique PRs")

    detail_columns = [
        "repo_id",
        "repo_full_name",
        "number",
        "id",
        "title",
        "body",
        "user",
        "user_type",
        "state",
        "created_at",
        "closed_at",
        "merged_at",
        "html_url",
        "additions",
        "deletions",
        "changed_files",
        "filenames",
        "commit_messages",
    ]
    repo_ids = selected["repo_id"].drop_duplicates().tolist()
    details = pq.read_table(
        task_path,
        columns=detail_columns,
        filters=[("repo_id", "in", repo_ids)],
    ).to_pandas()
    details["population_key"] = key_frame(details)
    details = details[details["population_key"].isin(selected["population_key"])]
    details = details.drop_duplicates("population_key")

    selected = selected.drop(
        columns=["repo_full_name", "html_url"], errors="ignore"
    ).merge(details, on=["repo_id", "number", "population_key"], validate="one_to_one")
    if len(selected) != 100:
        raise AssertionError(f"Expected 100 detailed rows, found {len(selected)}")

    selected["review_hash"] = [
        stable_hash(SEED, "review-order", repo_id, number)
        for repo_id, number in zip(selected["repo_id"], selected["number"], strict=True)
    ]
    selected = selected.sort_values("review_hash", kind="stable").reset_index(drop=True)
    selected["case_id"] = [f"LUNA-AUDIT-{index:03d}" for index in range(1, 101)]

    strata = (
        selected[
            [
                "audit_stratum",
                "stratum_population",
                "stratum_sample_size",
                "inclusion_probability",
                "sampling_weight",
            ]
        ]
        .drop_duplicates()
        .sort_values("audit_stratum")
        .reset_index(drop=True)
    )
    return selected, strata


THIN_GRAY = Side(style="thin", color="D9E2F3")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
DONE_FILL = PatternFill("solid", fgColor="E2F0D9")


def style_header(row) -> None:
    for cell in row:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN_GRAY)


def add_table(ws, name: str) -> None:
    ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def write_start_sheet(wb: Workbook, source_revision: str) -> None:
    ws = wb.active
    ws.title = "START_HERE"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "GPT-Luna performance-classification manual audit"
    ws["A1"].font = Font(size=18, bold=True, color="1F4E78")
    ws.merge_cells("A1:F1")

    rows = [
        ("Purpose", "Manually assess whether GPT-Luna correctly classified pull requests as performance or non-performance using the complete PR as evidence."),
        ("Unique PRs", 100),
        ("Design", "30 Luna-positive PRs from the final study sample, 30 Luna-positive PRs outside it, and 40 Luna-negative PRs."),
        ("Reviewer", "One blinded reviewer. Do not inspect the hidden METADATA sheet until all labels are complete."),
        ("Source revision", source_revision),
        ("Sampling seed", SEED),
        ("Labeled", '=COUNTIF(REVIEW!K2:K101,"<>")'),
        ("Performance", '=COUNTIF(REVIEW!K2:K101,"performance")'),
        ("Not performance", '=COUNTIF(REVIEW!K2:K101,"not_performance")'),
        ("Uncertain", '=COUNTIF(REVIEW!K2:K101,"uncertain")'),
    ]
    for row_index, (label, value) in enumerate(rows, start=3):
        ws.cell(row_index, 1, label).font = Font(bold=True)
        ws.cell(row_index, 1).fill = SUBHEADER_FILL
        ws.cell(row_index, 2, value)
        ws.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=6)
        ws.cell(row_index, 2).alignment = Alignment(wrap_text=True, vertical="top")

    ws["A15"] = "Workflow"
    ws["A15"].font = Font(size=14, bold=True, color="1F4E78")
    instructions = [
        "Read CODEBOOK before labeling.",
        "Work only in REVIEW. The queue is randomized and hides Luna's prediction, confidence, reason, author arm, and sampling stratum.",
        "Open the PR URL and inspect the complete PR, including diff and commits. The embedded body and file/commit lists are navigation aids.",
        "Select one label in the yellow manual_label column. Use observations only when you want to record an optional note.",
        "After all 100 cases are complete, wait at least seven days, unhide RECHECK, and label those 10 cases again without consulting REVIEW.",
        "After labeling, open RESULTS for the weighted confusion matrix and metrics. Only then unhide METADATA if case-level inspection is needed.",
    ]
    for index, text in enumerate(instructions, start=1):
        ws.cell(15 + index, 1, index)
        ws.cell(15 + index, 2, text)
        ws.merge_cells(start_row=15 + index, start_column=2, end_row=15 + index, end_column=6)
        ws.cell(15 + index, 2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 22
    for column in "BCDEF":
        ws.column_dimensions[column].width = 20


REVIEW_HEADERS = [
    "case_id",
    "pr_url",
    "repository",
    "pr_number",
    "title",
    "created_at",
    "change_size",
    "body",
    "files_changed",
    "commit_messages",
    "manual_label",
    "observations",
]


def review_values(row: pd.Series, case_id: str | None = None) -> list[object]:
    change_size = (
        f"+{int(row['additions'] or 0)} / -{int(row['deletions'] or 0)}; "
        f"{int(row['changed_files'] or 0)} files"
    )
    created_at = row["created_at"]
    if hasattr(created_at, "isoformat"):
        created_at = created_at.isoformat()
    return [
        case_id or row["case_id"],
        clean_cell(row["html_url"], 1_000),
        clean_cell(row["repo_full_name"], 1_000),
        int(row["number"]),
        clean_cell(row["title"], 5_000),
        clean_cell(created_at, 100),
        change_size,
        clean_cell(row["body"], 30_000),
        clean_cell(row["filenames"], 20_000),
        clean_cell(row["commit_messages"], 20_000),
        "",
        "",
    ]


def configure_review_sheet(ws, row_count: int, table_name: str) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{row_count + 1}"
    ws.sheet_view.showGridLines = False
    style_header(ws[1])
    add_table(ws, table_name)
    widths = {
        "A": 17,
        "B": 46,
        "C": 30,
        "D": 10,
        "E": 48,
        "F": 20,
        "G": 20,
        "H": 80,
        "I": 45,
        "J": 55,
        "K": 20,
        "L": 48,
    }
    for column, width in widths.items():
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=2, max_row=row_count + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in range(11, 13):
            row[column - 1].fill = INPUT_FILL
        url_cell = row[1]
        url_cell.hyperlink = str(url_cell.value).lstrip("'")
        url_cell.style = "Hyperlink"
    ws.conditional_formatting.add(
        f"A2:L{row_count + 1}",
        FormulaRule(formula=['$K2<>""'], fill=DONE_FILL),
    )

    validations = [("K", "=Lists!$A$2:$A$4")]
    for column, formula in validations:
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        validation.error = "Choose a value from the dropdown list."
        validation.errorTitle = "Invalid value"
        validation.prompt = "Use the predefined audit value."
        validation.promptTitle = "Manual audit"
        ws.add_data_validation(validation)
        validation.add(f"{column}2:{column}{row_count + 1}")


def write_review_sheet(wb: Workbook, sample: pd.DataFrame) -> None:
    ws = wb.create_sheet("REVIEW")
    ws.append(REVIEW_HEADERS)
    for _, row in sample.iterrows():
        ws.append(review_values(row))
    configure_review_sheet(ws, len(sample), "ReviewQueue")


def write_codebook(wb: Workbook) -> None:
    ws = wb.create_sheet("CODEBOOK")
    ws.sheet_view.showGridLines = False
    ws.append(["Field", "Value", "Operational definition"])
    rows = [
        ("manual_label", "performance", "The implemented change directly intends to improve runtime, latency, throughput, memory/allocation behavior, CPU/compute work, I/O/network efficiency, scalability/concurrency, build time, artifact size, or energy/cost."),
        ("manual_label", "not_performance", "The complete PR does not implement a performance improvement. Performance-related words, benchmark infrastructure, measurements, configuration, or dependency notes alone are insufficient."),
        ("manual_label", "uncertain", "The complete PR does not provide enough evidence for a defensible binary decision."),
        ("include", "Optimization code", "Include code changes that reduce work, allocations, latency, I/O, build work, or resource use even when no benchmark is reported."),
        ("exclude", "Benchmark/profiling only", "Exclude a PR that only adds or changes performance measurement infrastructure without implementing a performance improvement."),
        ("exclude", "Configuration only", "Exclude limits, timeouts, cache sizes, CI settings, or tuning parameters unless the PR implements a clear performance improvement as its task."),
        ("exclude", "Dependency update only", "Exclude routine dependency updates even if upstream release notes mention performance, unless this PR explicitly targets that improvement."),
        ("exclude", "Correctness/feature/refactor", "Exclude fixes, features, and refactorings with no explicit or technically evident performance objective."),
        ("observations", "Optional", "Use the observations field only when a short note will help explain an ambiguous or notable case."),
        ("scope", "Complete PR", "Use title, body, diff, files, commits, and relevant discussion. This is a construct-validity audit with richer evidence than Luna received."),
    ]
    for row in rows:
        ws.append(row)
    style_header(ws[1])
    add_table(ws, "Codebook")
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 110
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def write_results_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("RESULTS")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Automatic audit results"
    ws["A1"].font = Font(size=16, bold=True, color="1F4E78")
    ws.merge_cells("A1:D1")
    ws["A3"] = "Progress"
    ws["B3"] = '=COUNTIF(REVIEW!$K$2:$K$101,"<>")'
    ws["C3"] = "of 100"
    ws["A4"] = "Decided labels"
    ws["B4"] = '=COUNTIF(REVIEW!$K$2:$K$101,"performance")+COUNTIF(REVIEW!$K$2:$K$101,"not_performance")'
    ws["A5"] = "Uncertain labels"
    ws["B5"] = '=COUNTIF(REVIEW!$K$2:$K$101,"uncertain")'

    ws["A7"] = "Confusion component"
    ws["B7"] = "Raw sample count"
    ws["C7"] = "Weighted population estimate"
    style_header(ws[7])
    components = [
        (
            "True positive",
            '=COUNTIFS(METADATA!$I$2:$I$101,"perf",REVIEW!$K$2:$K$101,"performance")',
            '=SUMPRODUCT((METADATA!$I$2:$I$101="perf")*(REVIEW!$K$2:$K$101="performance")*METADATA!$R$2:$R$101)',
        ),
        (
            "False positive",
            '=COUNTIFS(METADATA!$I$2:$I$101,"perf",REVIEW!$K$2:$K$101,"not_performance")',
            '=SUMPRODUCT((METADATA!$I$2:$I$101="perf")*(REVIEW!$K$2:$K$101="not_performance")*METADATA!$R$2:$R$101)',
        ),
        (
            "False negative",
            '=COUNTIFS(METADATA!$I$2:$I$101,"<>perf",REVIEW!$K$2:$K$101,"performance")',
            '=SUMPRODUCT((METADATA!$I$2:$I$101<>"perf")*(REVIEW!$K$2:$K$101="performance")*METADATA!$R$2:$R$101)',
        ),
        (
            "True negative",
            '=COUNTIFS(METADATA!$I$2:$I$101,"<>perf",REVIEW!$K$2:$K$101,"not_performance")',
            '=SUMPRODUCT((METADATA!$I$2:$I$101<>"perf")*(REVIEW!$K$2:$K$101="not_performance")*METADATA!$R$2:$R$101)',
        ),
    ]
    for row_index, values in enumerate(components, start=8):
        for column, value in enumerate(values, start=1):
            ws.cell(row_index, column, value)

    ws["A14"] = "Metric"
    ws["B14"] = "Weighted estimate"
    style_header(ws[14])
    metrics = [
        ("Precision, all Luna positives", "=IFERROR(C8/(C8+C9),\"\")"),
        ("Recall, full Luna classification", "=IFERROR(C8/(C8+C10),\"\")"),
        ("Specificity", "=IFERROR(C11/(C11+C9),\"\")"),
        ("F1", "=IFERROR(2*B15*B16/(B15+B16),\"\")"),
        (
            "Precision, final study sample",
            '=IFERROR(SUMPRODUCT((METADATA!$H$2:$H$101=TRUE)*(REVIEW!$K$2:$K$101="performance")*METADATA!$R$2:$R$101)/SUMPRODUCT((METADATA!$H$2:$H$101=TRUE)*((REVIEW!$K$2:$K$101="performance")+(REVIEW!$K$2:$K$101="not_performance"))*METADATA!$R$2:$R$101),"")',
        ),
    ]
    for row_index, (label, formula) in enumerate(metrics, start=15):
        ws.cell(row_index, 1, label)
        ws.cell(row_index, 2, formula)
        ws.cell(row_index, 2).number_format = "0.0%"

    ws["A22"] = "Interpretation warnings"
    ws["A22"].font = Font(size=13, bold=True, color="1F4E78")
    warnings = [
        "Do not interpret metrics until all 100 cases are labeled.",
        "Weighted estimates target the auditable Luna population represented by the defined strata; deleted repositories and rows without a URL were excluded.",
        "Uncertain labels are excluded from metric denominators. Report best/worst-case sensitivity analyses separately.",
        "The negative sample is small for rare false negatives; recall may be unstable even after weighting.",
        "No confidence interval is calculated automatically because this is a disproportionate stratified design. Use the stratum information in SAMPLE_SUMMARY for design-aware uncertainty estimation.",
        "One reviewer provides no inter-rater reliability estimate. Use RECHECK only for intra-rater consistency.",
    ]
    for index, text in enumerate(warnings, start=1):
        ws.cell(22 + index, 1, index)
        ws.cell(22 + index, 2, text)
        ws.merge_cells(start_row=22 + index, start_column=2, end_row=22 + index, end_column=4)
        ws.cell(22 + index, 2).alignment = Alignment(wrap_text=True, vertical="top")

    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 31
    ws.column_dimensions["D"].width = 20
    for row in range(8, 12):
        ws.cell(row, 3).number_format = "0.0"


def write_summary_sheet(
    wb: Workbook, strata: pd.DataFrame, sample: pd.DataFrame, source_revision: str
) -> None:
    ws = wb.create_sheet("SAMPLE_SUMMARY")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Sampling design"
    ws["A1"].font = Font(size=16, bold=True, color="1F4E78")
    notes = [
        ("Source revision", source_revision),
        ("Seed", SEED),
        ("Unique PRs", len(sample)),
        ("Scope", "Only PRs classified by the LLM method and attributed as agentic or human-candidate; title_regex decisions and unresolved attribution are excluded."),
        ("Availability", "Deleted repositories and rows without a PR URL are excluded because complete manual review is impossible."),
        ("Inference", "Use stratum weights for population estimates. Unweighted totals describe only this deliberately balanced audit sample."),
        ("Recall caveat", "The 40 negatives provide an exploratory false-negative audit. A rare false-negative rate cannot be estimated tightly with this sample size."),
        ("Reviewer caveat", "One reviewer cannot establish inter-rater reliability. RECHECK provides an intra-rater consistency check."),
    ]
    for index, (label, value) in enumerate(notes, start=3):
        ws.cell(index, 1, label).font = Font(bold=True)
        ws.cell(index, 1).fill = SUBHEADER_FILL
        ws.cell(index, 2, value)
        ws.merge_cells(start_row=index, start_column=2, end_row=index, end_column=6)
        ws.cell(index, 2).alignment = Alignment(wrap_text=True, vertical="top")

    start = 13
    headers = list(strata.columns)
    for column, header in enumerate(headers, start=1):
        ws.cell(start, column, header)
    for row_index, values in enumerate(strata.itertuples(index=False), start=start + 1):
        for column, value in enumerate(values, start=1):
            ws.cell(row_index, column, value)
    style_header(ws[start])
    table = Table(
        displayName="SamplingStrata",
        ref=f"A{start}:{get_column_letter(len(headers))}{start + len(strata)}",
    )
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.column_dimensions["A"].width = 40
    for column in range(2, 7):
        ws.column_dimensions[get_column_letter(column)].width = 24


def write_metadata_sheet(wb: Workbook, sample: pd.DataFrame) -> None:
    ws = wb.create_sheet("METADATA")
    headers = [
        "case_id",
        "population_key",
        "repo_id",
        "repo_full_name",
        "number",
        "audit_stratum",
        "audit_arm",
        "in_final_sample",
        "luna_output",
        "luna_confidence",
        "luna_reason",
        "luna_model",
        "luna_classifier",
        "negative_label_group",
        "stratum_population",
        "stratum_sample_size",
        "inclusion_probability",
        "sampling_weight",
        "selection_hash",
        "review_hash",
    ]
    ws.append(headers)
    for _, row in sample.iterrows():
        ws.append(
            [
                row["case_id"],
                row["population_key"],
                row["repo_id"],
                row["repo_full_name"],
                row["number"],
                row["audit_stratum"],
                row["audit_arm"],
                bool(row["in_final_sample"]),
                row["aidev_task_type"],
                row["aidev_task_type_confidence"],
                clean_cell(row["aidev_task_type_reason"], 10_000),
                row["aidev_task_type_model"],
                row["aidev_task_type_classifier"],
                row["negative_label_group"],
                row["stratum_population"],
                row["stratum_sample_size"],
                row["inclusion_probability"],
                row["sampling_weight"],
                row["selection_hash"],
                row["review_hash"],
            ]
        )
    style_header(ws[1])
    add_table(ws, "AuditMetadata")
    ws.freeze_panes = "A2"
    ws.sheet_state = "hidden"


def write_recheck_sheet(wb: Workbook, sample: pd.DataFrame) -> None:
    recheck = sample.copy()
    recheck["recheck_hash"] = [
        stable_hash(SEED, "recheck", repo_id, number)
        for repo_id, number in zip(recheck["repo_id"], recheck["number"], strict=True)
    ]
    recheck = recheck.sort_values("recheck_hash", kind="stable").head(10)
    recheck = recheck.sample(frac=1, random_state=20260924).reset_index(drop=True)
    ws = wb.create_sheet("RECHECK")
    ws.append(REVIEW_HEADERS)
    for index, (_, row) in enumerate(recheck.iterrows(), start=1):
        ws.append(review_values(row, case_id=f"RECHECK-{index:02d}"))
    configure_review_sheet(ws, len(recheck), "RecheckQueue")
    ws.sheet_state = "hidden"


def write_lists_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Lists")
    columns = {
        "manual_label": ["performance", "not_performance", "uncertain"],
    }
    for column_index, (header, values) in enumerate(columns.items(), start=1):
        ws.cell(1, column_index, header)
        for row_index, value in enumerate(values, start=2):
            ws.cell(row_index, column_index, value)
    ws.sheet_state = "veryHidden"


def build_workbook(
    sample: pd.DataFrame, strata: pd.DataFrame, output: Path, source_revision: str
) -> None:
    wb = Workbook()
    wb.properties.title = "GPT-Luna performance classification manual audit"
    wb.properties.subject = "Blinded 100-PR precision and recall pilot"
    wb.properties.creator = "EMSE performance PR study"
    write_start_sheet(wb, source_revision)
    write_review_sheet(wb, sample)
    write_results_sheet(wb)
    write_codebook(wb)
    write_summary_sheet(wb, strata, sample, source_revision)
    write_metadata_sheet(wb, sample)
    write_recheck_sheet(wb, sample)
    write_lists_sheet(wb)
    wb.active = wb.sheetnames.index("START_HERE")
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

    # Reopen immediately to catch malformed XML, broken tables, or invalid visibility.
    checked = load_workbook(output, data_only=False)
    if checked["REVIEW"].max_row != 101:
        raise AssertionError("Workbook validation failed: REVIEW must contain 100 cases")
    if checked["METADATA"].sheet_state != "hidden":
        raise AssertionError("Workbook validation failed: METADATA must remain hidden")
    checked.close()


def main() -> None:
    args = parse_args()
    sample, strata = build_sample(
        args.curated_labels, args.task_decisions, args.balanced_sample
    )
    build_workbook(sample, strata, args.output, args.source_revision)
    print(f"Wrote {args.output}")
    print(f"Cases: {len(sample)}")
    print(sample["audit_stratum"].value_counts().sort_index().to_string())


if __name__ == "__main__":
    main()
